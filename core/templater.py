# -*- coding: utf-8 -*-
"""模板整合输出（F17）：把比价/匹配结果按上级给的模板 xlsx 写入指定区域。

思路（复用对账生成器「模板+数据」）：模板里预留抬头/表头样式，数据从指定起始单元格
整块写入（含表头行）；保留模板原有格式与其它内容。仅支持 xlsx。
"""
import io

from .registry import skill


@skill(
    name="模板整合输出",
    desc="把结果矩阵写入模板 xlsx 的指定起始单元格（保留模板格式），导出填充后的文件",
    inputs={"template_bytes": "模板文件字节", "matrix_df": "结果矩阵",
            "sheet": "写入的目标 Sheet（默认第一个）", "start_cell": "起始单元格，如 A3",
            "title": "可选：写入首行的标题"},
    outputs={"bytes": "填充后的 xlsx 字节", "info": "写入位置信息"},
    task_modes=["完整比价", "两表匹配补缺"],
)
def fill_template(template_bytes, matrix_df, sheet=None, start_cell="A3",
                  include_header=True, title=None):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import coordinate_to_tuple

    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb[sheet] if (sheet and sheet in wb.sheetnames) else wb[wb.sheetnames[0]]
    row0, col0 = coordinate_to_tuple(start_cell or "A3")
    r = row0
    if title:
        ws.cell(row=r, column=col0, value=title)
        r += 1
    cols = list(matrix_df.columns)
    if include_header:
        for j, c in enumerate(cols):
            ws.cell(row=r, column=col0 + j, value=str(c))
        r += 1
    for _, row in matrix_df.iterrows():
        for j, c in enumerate(cols):
            v = row[c]
            ws.cell(row=r, column=col0 + j, value=("" if v is None else v))
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue(), {"sheet": ws.title, "start": start_cell,
                            "rows": len(matrix_df), "cols": len(cols),
                            "end": f"{get_column_letter(col0 + len(cols) - 1)}{r - 1}"}
