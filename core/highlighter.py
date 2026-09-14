# -*- coding: utf-8 -*-
"""最低价标红：矩阵内逐行找最低价并高亮（独立技能；网页展示 + Excel 落色两用）。"""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .converter import to_number
from .registry import skill

# 浅色彩虹预设（名称, 底色, 字色）——默认「浅黄」（WPS 习惯）
PRESETS = [
    ("浅黄", "#FFEB9C", "#9C6500"),
    ("浅红", "#FFC7CE", "#9C0006"),
    ("浅橙", "#FFD9B3", "#9C5700"),
    ("浅绿", "#C6EFCE", "#006100"),
    ("浅青", "#CCECEF", "#00666E"),
    ("浅蓝", "#DDEBF7", "#1F4E79"),
    ("浅紫", "#E4DFEC", "#5B2C6F"),
    ("浅粉", "#FCE4EC", "#9C1E5A"),
]
DEFAULT_FILL, DEFAULT_FONT = "#FFEB9C", "#9C6500"


def _argb(hex_color, default="FFFFFFFF"):
    """#RRGGBB / RRGGBB → ARGB 8 位（openpyxl 需要）。"""
    s = str(hex_color or "").strip().lstrip("#")
    if len(s) == 6:
        return "FF" + s.upper()
    if len(s) == 8:
        return s.upper()
    return default


def _fill_font(fill_hex, font_hex):
    return (PatternFill("solid", fgColor=_argb(fill_hex, "FFFFEB9C")),
            Font(color=_argb(font_hex, "FF9C6500"), bold=True))


@skill(
    name="最低价标红",
    desc="比价矩阵逐行找最低价；返回每行最低值与供应商，并可写入 Excel 高亮",
    inputs={"df": "矩阵 DataFrame(行=品类)", "supplier_cols": "供应商价格列名列表"},
    outputs={"min_vals": "每行最低价", "min_sups": "每行最低价供应商列名"},
    task_modes=["完整比价"],
)
def find_min(df, supplier_cols):
    min_vals, min_sups = [], []
    for _, row in df[supplier_cols].iterrows():
        nums = {c: to_number(v) for c, v in row.items()}
        nums = {c: n for c, n in nums.items() if n is not None and n > 0}
        if not nums:
            min_vals.append("")
            min_sups.append("")
        else:
            best = min(nums, key=nums.get)
            min_vals.append(nums[best])
            min_sups.append(best)
    return min_vals, min_sups


def export_highlighted(matrix_df, supplier_cols, out_path,
                       min_col="最低价", sup_col="最低价供应商",
                       min_fill=DEFAULT_FILL, min_font=DEFAULT_FONT):
    """把比价矩阵写为 xlsx：每行最低价单元格标红，附最低价供应商两列。

    min_fill / min_font：#RRGGBB 或 ARGB；配色见 PRESETS（默认浅黄）。
    """
    from openpyxl import Workbook

    min_vals, min_sups = find_min(matrix_df, supplier_cols)
    fill, font = _fill_font(min_fill, min_font)
    out = matrix_df.copy()
    out[min_col] = min_vals
    out[sup_col] = min_sups

    wb = Workbook()
    ws = wb.active
    ws.title = "比价矩阵"
    cols = list(out.columns)
    ws.append([str(c) for c in cols])
    for row in out.itertuples(index=False, name=None):
        ws.append(["" if _is_na(v) else v for v in row])

    eps = 1e-9
    for r in range(2, ws.max_row + 1):
        mv = min_vals[r - 2]
        if mv == "":
            continue
        for c_idx, col in enumerate(cols, start=1):
            if col in supplier_cols:
                n = to_number(ws.cell(row=r, column=c_idx).value)
                if n is not None and abs(n - mv) < eps:
                    cell = ws.cell(row=r, column=c_idx)
                    cell.fill = fill
                    cell.font = font
                    cell.alignment = Alignment(horizontal="center")
    for c_idx in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 18
    ws.freeze_panes = "A2"
    wb.save(out_path)
    return out_path


def _is_na(v):
    try:
        import pandas as pd
        return pd.isna(v)
    except Exception:
        return v is None
