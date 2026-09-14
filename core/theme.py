# -*- coding: utf-8 -*-
"""表格美化：统一主题（表头/列宽/冻结/筛选/数字格式/语义色/页脚），业务导出与独立美化共用。

配色（低饱和浅底+深字，不遮数据；语义互不冲突）：
- 结论/推荐 → 浅绿；风险/异常 → 浅红；缺失/未匹配 → 浅灰；置信高/中/低 → 浅黄/浅蓝/浅紫
"""
import datetime as _dt
import io
import os

from .registry import skill

# ---------- 语义色板（静态填充） ----------
COLORS = {
    "conclusion": ("#E2EFDA", "#375623"),   # 结论/推荐（最低价）
    "risk":       ("#FCE4E4", "#9C0006"),   # 风险/异常
    "missing":    ("#F2F2F2", None),        # 缺失/未匹配（留空）
    "conf_high":  ("#FFF2CC", "#7F6000"),   # 置信 80–99%
    "conf_mid":   ("#DDEBF7", "#1F4E79"),   # 置信 40–80%
    "conf_low":   ("#E4DFEC", "#5B2C6F"),   # 置信 20–40%
}
HEADER_FILL, HEADER_FONT = "#44546A", "#FFFFFF"

_MONEY_KW = ("金额", "价", "工资", "费用", "成本", "总价", "元", "薪")
_PCT_KW = ("率", "比例", "占比")
_DATE_KW = ("日期", "时间", "年月")
# 最低值高亮时排除这些"不可比"的数值列（序号/时间/数量等）
_MIN_SKIP_KW = ("序号", "编号", "编码", "日期", "时间", "年月", "月", "天", "年", "数量",
                "单位", "电话", "手机", "岗位数", "人数", "月数")


def _argb(hex_color):
    return "FF" + str(hex_color).lstrip("#").upper()


def _disp_width(s):
    return sum(2 if ord(ch) > 127 else 1 for ch in str(s))


def estimate_width(values, header="", min_w=8, max_w=40):
    """按内容估算列宽（中文按 2 字符）；min/max 夹紧。"""
    w = _disp_width(header) + 2
    for v in list(values)[:200]:
        if v is None:
            continue
        w = max(w, _disp_width(v))
    return int(min(max(w + 2, min_w), max_w))


def guess_number_format(header, values):
    """按列头 + 取值推断数字格式；无法判断返回 None（保持原样）。"""
    hs = str(header or "")
    vals = [v for v in values if v is not None and not (isinstance(v, str) and not v.strip())]
    if not vals:
        return None
    if all(isinstance(v, (_dt.datetime, _dt.date)) for v in vals):
        return "yyyy-mm-dd"
    if any(k in hs for k in _DATE_KW) and any(isinstance(v, (_dt.datetime, _dt.date)) for v in vals):
        return "yyyy-mm-dd"
    nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if len(nums) != len(vals):
        return None
    if "%" in hs or any(k in hs for k in _PCT_KW):
        return "0.00%" if max(abs(float(v)) for v in nums) <= 1 else '0.0"%"'
    if all(float(v).is_integer() for v in nums):
        return "#,##0"
    return "#,##0.00"


def style_sheet(ws, header_row=1, *, autosize=True, freeze=True, autofilter=True,
                number_formats=True, banded=False, highlight_min=False,
                footer_lines=None, min_w=8, max_w=40, skip_autosize_rows=5000):
    """对已有工作表做统一美化（表头样式/列宽/冻结/筛选/数字格式/可选最低值高亮/页脚）。"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    hr = int(header_row)
    ncols = ws.max_column
    nrows = ws.max_row
    # 表头
    for c in range(1, ncols + 1):
        cell = ws.cell(row=hr, column=c)
        cell.font = Font(bold=True, color=_argb(HEADER_FONT))
        cell.fill = PatternFill("solid", fgColor=_argb(HEADER_FILL))
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[hr].height = 24
    # 数据区：默认**居中**（表头也居中）；斑马纹
    for r in range(hr + 1, nrows + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if banded and (r - hr) % 2 == 0 and cell.fill.patternType is None:
                cell.fill = PatternFill("solid", fgColor="FFF7F9FC")
    # 数字格式
    if number_formats:
        for c in range(1, ncols + 1):
            fmt = guess_number_format(ws.cell(row=hr, column=c).value,
                                      [ws.cell(row=r, column=c).value for r in range(hr + 1, nrows + 1)])
            if fmt:
                for r in range(hr + 1, nrows + 1):
                    if ws.cell(row=r, column=c).value is not None:
                        ws.cell(row=r, column=c).number_format = fmt
    # 最低值高亮（可选）
    if highlight_min:
        fill, font = COLORS["conclusion"]
        min_cols = [c for c in range(1, ncols + 1)
                    if not any(k in str(ws.cell(row=hr, column=c).value or "") for k in _MIN_SKIP_KW)]
        for r in range(hr + 1, nrows + 1):
            nums = [(c, ws.cell(row=r, column=c).value) for c in min_cols
                    if isinstance(ws.cell(row=r, column=c).value, (int, float))
                    and not isinstance(ws.cell(row=r, column=c).value, bool)]
            if not nums:
                continue
            m = min(v for _, v in nums)
            for c, v in nums:
                if v == m:
                    cell = ws.cell(row=r, column=c)
                    cell.fill = PatternFill("solid", fgColor=_argb(fill))
                    cell.font = Font(bold=True, color=_argb(font))
    # 列宽
    if autosize:
        if nrows > skip_autosize_rows:
            for c in range(1, ncols + 1):
                ws.column_dimensions[get_column_letter(c)].width = 16
        else:
            for c in range(1, ncols + 1):
                vals = [ws.cell(row=r, column=c).value for r in range(hr, nrows + 1)]
                ws.column_dimensions[get_column_letter(c)].width = estimate_width(
                    vals[1:], vals[0], min_w, max_w)
    # 冻结 + 筛选
    if freeze:
        ws.freeze_panes = ws.cell(row=hr + 1, column=1).coordinate
    if autofilter and nrows > hr:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(ncols)}{nrows}"
    # 页脚
    if footer_lines:
        r = nrows + 2
        for line in footer_lines:
            ws.cell(row=r, column=1, value=str(line)).alignment = Alignment(horizontal="left")
            r += 1
    return ws


def footer_lines(source="", sheet="", extra=None):
    lines = [f"生成时间：{_dt.datetime.now():%Y-%m-%d %H:%M}"]
    if source:
        lines.append(f"来源：{source}" + (f" · Sheet「{sheet}」" if sheet else ""))
    lines += list(extra or [])
    return lines


@skill(
    name="表格美化",
    desc="统一美化 xlsx：表头深底白字/列宽自适应/冻结/筛选/数字格式/可选最低值高亮/页脚备注",
    inputs={"bytes": "xlsx 字节", "filename": "文件名", "sheet": "目标 Sheet",
            "header_row": "表头行(1-based)", "options": "autosize/freeze/autofilter/number_formats/banded/highlight_min/new_sheet"},
    outputs={"bytes": "美化后的 xlsx 字节", "info": "处理信息"},
    task_modes=["表格美化"],
)
def beautify_bytes(data, filename, sheet, header_row=1, options=None):
    """独立美化：默认**新增 sheet「美化-原名」**（原表不动）；options['in_place']=True 则就地美化（调用方负责先备份）。"""
    from openpyxl import load_workbook
    opt = dict(options or {})
    wb = load_workbook(io.BytesIO(data))
    if sheet not in wb.sheetnames:
        wb.close()
        raise ValueError(f"工作簿中找不到 Sheet「{sheet}」")
    if opt.get("in_place"):
        target = wb[sheet]
    else:
        target = wb.copy_worksheet(wb[sheet])
        base = f"美化-{sheet}"
        name, i = base, 0
        while name in wb.sheetnames:
            i += 1
            name = f"{base}({i})"
        target.title = name
    style_sheet(target, header_row,
                autosize=opt.get("autosize", True), freeze=opt.get("freeze", True),
                autofilter=opt.get("autofilter", True),
                number_formats=opt.get("number_formats", True),
                banded=opt.get("banded", False),
                highlight_min=opt.get("highlight_min", False),
                footer_lines=footer_lines(filename, sheet,
                                          ["（浅绿=各行最低值）"] if opt.get("highlight_min") else None))
    buf = io.BytesIO()
    wb.save(buf)
    info = {"sheet": target.title, "header_row": header_row, "in_place": bool(opt.get("in_place"))}
    wb.close()
    return buf.getvalue(), info


def beautify_file_in_place(path, sheet, header_row=1, options=None):
    """就地美化并写回原文件：**先建 ai副本 备份 → 重读含备份的最新内容 → 美化 → 写回**。

    关键：备份必须先落盘、再重读字节，否则用旧的"内存字节"写回会把刚建的备份覆盖掉（旧 bug）。
    返回 (backup_sheet_name, info)；文件被占用抛 PermissionError。
    """
    from .excel_io import backup_sheet_numbered
    opt = dict(options or {})
    opt["in_place"] = True
    bk = backup_sheet_numbered(path, sheet)          # ① 备份（写盘，已存在则 (1)(2) 递增）
    with open(path, "rb") as f:                      # ② 重读"含备份"的最新内容
        fresh = f.read()
    out, info = beautify_bytes(fresh, os.path.basename(path), sheet, header_row, opt)
    try:
        with open(path, "wb") as f:                  # ③ 写回（备份与美化后的 sheet 同在工作簿内）
            f.write(out)
    except PermissionError:
        raise PermissionError(f"文件被占用（可能正在 WPS/Excel 中打开）：{path}")
    return bk, info
