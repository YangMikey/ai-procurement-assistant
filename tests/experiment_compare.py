# -*- coding: utf-8 -*-
"""M2c 完整比价端到端：解析 → 对齐 → 标红导出（颜色可配）→ 采购建议。

- 合成 4 家跑全流程，核对矩阵/最低价/建议
- 单测：标红颜色参数（导出 xlsx 回读 fill 颜色）、异常低价、缺报价
- 报告：data/outputs/比价报告.txt
运行：py tests/experiment_compare.py
"""
import json
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

import pandas as pd
from openpyxl import load_workbook

from core.advisor_rule import advise
from core.aligner import align_quotes
from core.highlighter import PRESETS, export_highlighted, find_min
from core.parser_rule import parse_quote_file

SAMPLES = os.path.join(_ROOT, "data", "samples")
OUT = os.path.join(_ROOT, "data", "outputs")
TMP = os.path.join(os.environ["TEMP"], "opencode", "cmp_demo.xlsx")
os.makedirs(os.path.dirname(TMP), exist_ok=True)

ok = 0
lines = []


def check(name, cond):
    global ok
    assert cond, f"[FAIL] {name}"
    ok += 1
    print(f"[PASS] {name}")
    lines.append(f"[PASS] {name}")


# ---- 全流程 ----
paths = [os.path.join(SAMPLES, f) for f in [
    "报价单A_标准格式.xlsx", "报价单B_表头第2行_仅不含税.xlsx",
    "报价单C_文本价格.xlsx", "报价单D_同品异名_缺行多行.xlsx"]]
quotes = [parse_quote_file(path=p)[0] for p in paths]
res = align_quotes(quotes, price_field="不含税单价")
m, sup = res["matrix"], res["suppliers"]
check("比价：矩阵 7 行 × 4 供应商", len(m) == 7 and all(s in m.columns for s in sup))
check("比价：含最低价/最低价供应商列", {"最低价", "最低价供应商"} <= set(m.columns))
check("比价：含「对齐备注」列（来源可查）", "对齐备注" in m.columns)

# ---- 标红颜色可配（自定义浅绿）→ 回读校验 ----
export_highlighted(m, sup, TMP, min_fill="#C6EFCE", min_font="#006100")
wb = load_workbook(TMP)
ws = wb.active
cols = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
i_sup = {s: cols.index(s) + 1 for s in sup}
i_name = cols.index("品名") + 1
i_min = cols.index("最低价供应商") + 1
r_screw = next(r for r in range(2, ws.max_row + 1) if "螺钉" in str(ws.cell(row=r, column=i_name).value))
best = ws.cell(row=r_screw, column=i_min).value
cell_min = ws.cell(row=r_screw, column=i_sup[best])
check("标红：最低价单元格 fill=自定义浅绿(FFC6EFCE)", str(cell_min.fill.fgColor.rgb).upper() == "FFC6EFCE")
check("标红：最低价单元格字色=深绿(FF006100)", str(cell_min.font.color.rgb).upper() == "FF006100")
wb.close()
check("调色板：含浅色彩虹 ≥8 档", len(PRESETS) >= 8 and any(p[0] == "浅黄" for p in PRESETS))

# ---- 采购建议：缺报价 + 异常低价 ----
adv = advise(m, sup)
check("建议：识别到缺报价品类（D 缺 LED）",
      any("LED" in str(x) for x in adv["rows"][adv["rows"]["报价数"] < len(sup)]["品名"]))
check("建议：汇总含异常低价/缺报价计数",
      "异常低价" in adv["summary"] and "有缺报价的品类" in adv["summary"])

_m2 = pd.DataFrame({"品名": ["A品", "B品"], "规格": ["", ""],
                    "甲": [100.0, 10.0], "乙": [110.0, 11.0], "丙": [5.0, 12.0]})
adv2 = advise(_m2, ["甲", "乙", "丙"], low_ratio=0.5)
check("建议：异常低价被标记（5 远低于中位数 100）",
      adv2["rows"].iloc[0]["预警"] == "异常低价" and adv2["summary"]["异常低价"] == 1)
check("建议：正常行不误报", adv2["rows"].iloc[1]["预警"] == "")

# ---- 生成报告 ----
os.makedirs(OUT, exist_ok=True)
with open(os.path.join(OUT, "比价报告.txt"), "w", encoding="utf-8") as f:
    f.write("=" * 72 + "\n完整比价报告（M2c，合成 4 家）\n" + "-" * 72 + "\n")
    f.write(f"品类数={len(m)} 供应商={sup}\n")
    f.write("最低价供应商：" + "；".join(
        f"{r['品名']}→{r['最低价供应商']}" for _, r in m.iterrows()) + "\n")
    f.write(f"建议汇总：{adv['summary']}\n")
    for t in adv["text"]:
        f.write(t + "\n")
    f.write("=" * 72 + "\n")
os.remove(TMP)

print(f"\n===== M2c 完整比价端到端通过：{ok} 项断言；报告见 data\\outputs\\比价报告.txt =====")
