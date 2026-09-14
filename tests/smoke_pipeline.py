# -*- coding: utf-8 -*-
"""端到端冒烟测试：两表匹配 → 合并比价矩阵 → 最低价标红导出 → 换算回环。

运行：py tests/smoke_pipeline.py
通过标准：全部断言通过 + 生成 data/samples/比价矩阵_测试.xlsx
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.converter import convert_tax, to_number
from core.excel_io import read_table
from core.highlighter import find_min, export_highlighted
from core.matcher import run_match

S = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "samples")
A_PATH = os.path.join(S, "报价单A_标准格式.xlsx")
C_PATH = os.path.join(S, "报价单C_文本价格.xlsx")
OUT_PATH = os.path.join(S, "比价矩阵_测试.xlsx")

ok = 0


def check(name, cond):
    global ok
    assert cond, f"[FAIL] {name}"
    ok += 1
    print(f"[PASS] {name}")


# ---- 1. 两表匹配（A 钥匙=品类 ↔ C 钥匙=品类，模糊模式处理文本价格） ----
a_df = read_table(A_PATH, "报价", header_row=1)          # 品类/规格型号/单位/数量/含税单价/税率
c_df = read_table(C_PATH, "C供应商报价", header_row=1)     # 品类/型号/单位/含税单价(元)
res = run_match(a_df, c_df, ["品类"], ["品类"], [], [], ["含税单价(元)"],
                mode="fuzzy", threshold=70)
check("匹配：6 行全部钥匙精确命中", res["stats"]["钥匙精确"] == 6)
check("匹配：无未匹配", res["stats"]["未匹配"] == 0)

# ---- 2. 合并比价矩阵（行=品类，列=A/C 价格） ----
a_price = dict(zip(a_df["品类"], a_df["含税单价"]))
c_price = dict(zip(c_df["品类"], c_df["含税单价(元)"]))
matrix = res["result"][["品类", "含税单价", "含税单价(元)"]].copy()
matrix = matrix.rename(columns={"含税单价": "供应商A", "含税单价(元)": "供应商C"})
check("矩阵：6 行 2 家供应商", len(matrix) == 6 and list(matrix.columns)[1:] == ["供应商A", "供应商C"])

# ---- 3. 最低价标红导出 ----
export_highlighted(matrix, ["供应商A", "供应商C"], OUT_PATH)
check("标红导出：文件生成", os.path.exists(OUT_PATH))
min_vals, min_sups = find_min(matrix, ["供应商A", "供应商C"])
check("标红：每行都有最低价供应商", all(s in ("供应商A", "供应商C") for s in min_sups))

# ---- 4. 换算回环（含税→不含税→含税 应复原；文本价'¥1,234.56'能解析） ----
test_df = a_df.copy()
out, added = convert_tax(test_df, ["含税单价"], tax_rate=0.13, direction="to_ex")
check("换算：新增列", added == ["含税单价(不含税)"])
back, _ = convert_tax(out, ["含税单价(不含税)"], tax_rate=0.13, direction="to_in")
pairs = zip(a_df["含税单价"].tolist(), back["含税单价(不含税)(含税)"].tolist())
check("换算：回环复原（误差<0.01）", all(abs(to_number(x) - to_number(y)) < 0.01 for x, y in pairs))
check("解析：文本价 '¥...' 可读", to_number(c_df["含税单价(元)"].iloc[0]) is not None)

# ---- 4.5 插入位置：新列必须在源列右侧 ----
check("插入位置：新列紧贴源列右侧",
      out.columns.get_loc("含税单价") + 1 == out.columns.get_loc("含税单价(不含税)"))

# ---- 4.6 每列独立税率（无配对）+ 列名自动识别 + 税率单位规则 ----
from core.converter import parse_rate, parse_rate_from_name
check("列名识别：'同国贸易 税率13%'→0.13", parse_rate_from_name("同国贸易 税率13%") == 0.13)
check("列名识别：'玉兰邮（1%）'→0.01", parse_rate_from_name("玉兰邮（1%）") == 0.01)
check("列名识别：'润锋 （13%）'→0.13", parse_rate_from_name("润锋 （13%）") == 0.13)
check("列名识别：'利源通 税率（1%）'→0.01", parse_rate_from_name("利源通 税率（1%）") == 0.01)
check("列名识别：真实表头'利源通 税率（1%)'无%号→0.01", parse_rate_from_name("利源通 税率（1%)") == 0.01)
check("列名识别：无标记列 → None", parse_rate_from_name("Unnamed: 8") is None)
check("税率单位：'1'→1%", parse_rate("1") == 0.01)
check("税率单位：'13'→13%", parse_rate("13") == 0.13)
check("税率单位：'13%'→13%", parse_rate("13%") == 0.13)
check("税率单位：'1%'→1%", parse_rate("1%") == 0.01)

import pandas as pd
two = pd.DataFrame({"A列 税率13%": [113.0], "B列（1%）": [101.0], "C列": [130.0]})
o2, a2 = convert_tax(two, ["A列 税率13%", "B列（1%）", "C列"], tax_rate=0.13,
                     direction="to_ex", rates=[0.13, 0.01, None])
check("独立税率：A列 113/1.13=100", abs(to_number(o2["A列 税率13%(不含税)"].iloc[0]) - 100.0) < 0.01)
check("独立税率：B列 101/1.01=100", abs(to_number(o2["B列（1%）(不含税)"].iloc[0]) - 100.0) < 0.01)
check("独立税率：C列未指定走全局 130/1.13≈115.04", abs(to_number(o2["C列(不含税)"].iloc[0]) - 115.0442) < 0.01)
check("独立税率：三列各自插在源列右侧",
      o2.columns.get_loc("A列 税率13%") + 1 == o2.columns.get_loc("A列 税率13%(不含税)")
      and o2.columns.get_loc("B列（1%）") + 1 == o2.columns.get_loc("B列（1%）(不含税)")
      and o2.columns.get_loc("C列") + 1 == o2.columns.get_loc("C列(不含税)"))

# ---- 4.7 ai副本递增备份 + 换算写回原文件 ----
import shutil
from core.excel_io import backup_sheet_numbered, writeback_convert

TMP = os.path.join(S, "_tmp_writeback_test.xlsx")
shutil.copyfile(A_PATH, TMP)
b1 = backup_sheet_numbered(TMP, "报价")
b2 = backup_sheet_numbered(TMP, "报价")
check("ai副本：首次命名为 ai副本报价", b1 == "ai副本报价")
check("ai副本：重复运行递增 (1)", b2 == "ai副本报价(1)")

conv_vals = out["含税单价(不含税)"].tolist()
written = writeback_convert(TMP, "报价", 1, [(5, "含税单价(不含税)", conv_vals)])
check("写回：返回写入列名", written == ["含税单价(不含税)"])
re_df = read_table(TMP, "报价", header_row=1)
check("写回：新列位于源列右侧",
      list(re_df.columns).index("含税单价") + 1 == list(re_df.columns).index("含税单价(不含税)"))
re_vals = re_df["含税单价(不含税)"].tolist()
check("写回：数值与预期一致",
      all(abs(to_number(a) - to_number(b)) < 0.01 for a, b in zip(re_vals, conv_vals)))
written2 = writeback_convert(TMP, "报价", 1, [(5, "含税单价(不含税)", conv_vals)])
check("幂等：重跑不重复插列（返回标注覆写）", written2 == ["(覆写)含税单价(不含税)"])
re_df2 = read_table(TMP, "报价", header_row=1)
check("幂等：重跑后新列仍只有一列",
      list(re_df2.columns).count("含税单价(不含税)") == 1)
os.remove(TMP)

# ---- 4.8 合计行跳过 ----
skip_df = pd.DataFrame({"品名": ["胶条", "合计"], "含税单价": [113.0, "SUM文字"]})
skip_out, _ = convert_tax(skip_df, ["含税单价"], tax_rate=0.13, direction="to_ex",
                          skip_mask=[False, True])
check("合计跳过：正常行已换算", abs(to_number(skip_out["含税单价(不含税)"].iloc[0]) - 100.0) < 0.01)
check("合计跳过：合计行输出空（写回时原值不动）", skip_out["含税单价(不含税)"].iloc[1] == "")

# ---- 4.9 两行式（合并）表头拆分 ----
from core.excel_io import split_two_row_header
TMP2 = os.path.join(S, "_tmp_two_row_header.xlsx")
from openpyxl import Workbook as _WB
_wb = _WB(); _ws = _wb.active; _ws.title = "报价"
_ws["F1"] = "同国贸易 税率13%"
_ws.merge_cells("F1:F2")
_ws.append([]); _ws.append([])
for c, h in enumerate(["序号", "品名", "规格", "单位", "数量"], start=1):
    _ws.cell(row=2, column=c, value=h)
_ws.append([1, "胶条", "黑", "卷", 1, 113.0])
_wb.save(TMP2)
with open(TMP2, "rb") as _fh:
    _b = _fh.read()
names, is_two = split_two_row_header(_b, TMP2, "报价", 2)
check("两行表头：检测为两行式", is_two is True)
check("两行表头：F列拆分为「同国贸易 税率13%」", "同国贸易 税率13%" in (names or []))
check("两行表头：左侧列保持原名", (names or [])[0] == "序号" and (names or [])[4] == "数量")

# ---- 4.10 顶行自动降级 + 横向合并前缀回填 + 数据区合并填充 ----
from core.excel_io import detect_header_block_bottom, unmerge_fill_data
TMP3 = os.path.join(S, "_tmp_merge_block.xlsx")
_wb = _WB(); _ws = _wb.active; _ws.title = "进度"
# 两行表头：A1:A2 事业部、B1:B2 项目、C1:C2 类别；D1:F1 合同信息（横向）+ 第二行月份
for c, h in enumerate(["事业部", "项目", "类别"], start=1):
    _ws.cell(row=1, column=c, value=h)
    _ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
_ws["D1"] = "合同信息"
_ws.merge_cells("D1:F1")
for c, m in enumerate(["2026年10月", "2026年11月", "2026年12月"], start=4):
    _ws.cell(row=2, column=c, value=m)
_ws["G1"] = "责任人"; _ws.merge_cells("G1:G2")
_ws["H1"] = "合同进度"; _ws.merge_cells("H1:H2")
# 数据区：类别竖向合并 C3:C4；责任人竖向合并 G3:G4
_ws.append(["一部", "朗晴居", None, None, None, None, None, "正常"])
_ws.append([None, "景安花园", None, None, None, None, None, "正常"])
_ws.merge_cells("C3:C4")
_ws.merge_cells("G3:G4")
_ws["C3"] = "日常保洁"; _ws["G3"] = "张三"
_wb.save(TMP3)
with open(TMP3, "rb") as _fh:
    _b3 = _fh.read()
check("顶行检测：第1行是两行表头顶部 → 底行=2", detect_header_block_bottom(_b3, TMP3, "进度", 1) == 2)
check("顶行检测：第2行（底行）不再降级", detect_header_block_bottom(_b3, TMP3, "进度", 2) is None)
names3, is_two3 = split_two_row_header(_b3, TMP3, "进度", 2)
check("前缀回填：横向合并下的月份带「合同信息」前缀",
      is_two3 and "合同信息 2026年11月" in names3 and "合同信息 2026年12月" in names3)
check("前缀回填：竖向合并列保持原名", is_two3 and names3[0] == "事业部" and names3[6] == "责任人")
m_df3_raw = read_table(TMP3, "进度", header_row=1)
check("顶行检测：直读第1行时列名有 Unnamed（证明降级必要）",
      any(str(c).startswith("Unnamed") for c in m_df3_raw.columns))
m_df3 = read_table(TMP3, "进度", header_row=2)
m_df3 = m_df3.rename(columns={m_df3.columns[j]: names3[j]
                              for j in range(min(len(names3), len(m_df3.columns)))})
filled, n_fill = unmerge_fill_data(_b3, TMP3, "进度", 2, m_df3)
check("合并填充：类别列每行都有值",
      to_number(filled["类别"].iloc[0]) is None and str(filled["类别"].iloc[0]) == "日常保洁"
      and str(filled["类别"].iloc[1]) == "日常保洁")
check("合并填充：责任人列每行都有值", str(filled["责任人"].iloc[0]) == "张三"
      and str(filled["责任人"].iloc[1]) == "张三")
check("合并填充：已填格数=2（合并覆盖的空格 C4/G4）", n_fill == 2)
os.remove(TMP3)
os.remove(TMP2)

# ---- 4.11 写回：两行合并表头下，别名指向已有列要「就地更新」而不是追加新列 ----
from core.excel_io import write_result_in_sheet
from openpyxl import load_workbook
TMP4 = os.path.join(S, "_tmp_alias_writeback.xlsx")
_wb = _WB(); _ws = _wb.active; _ws.title = "进度"
_ws["A1"] = "事业部"; _ws.merge_cells("A1:A2")   # 两行合并表头：标签在第1行
_ws["B1"] = "项目"; _ws.merge_cells("B1:B2")
_ws.append(["", "朗晴居"])    # 物理行3：事业部列为空（待补）
_ws.append(["", "景安花园"])  # 物理行4
_wb.save(TMP4)
_st = {0: {"status": "钥匙精确", "combos": [(0, 100.0)], "reason": ""},
       1: {"status": "钥匙精确", "combos": [(1, 100.0)], "reason": ""}}
_written, _overwritten = write_result_in_sheet(
    TMP4, "进度", 2, [0, 1], _st,
    [["华二第三事业部"], ["华二第一事业部"]], ["事业部"])
check("写回别名：写入列=事业部+3注释列，且无重复事业部",
      _written[0] == "事业部" and "事业部" not in _written[1:])
_wb2 = load_workbook(TMP4); _ws2 = _wb2["进度"]
check("写回别名：A列（合并表头已有列）被就地填充",
      _ws2.cell(row=3, column=1).value == "华二第三事业部"
      and _ws2.cell(row=4, column=1).value == "华二第一事业部")
check("写回别名：事业部列只出现一次（没有重复追加）",
      sum(1 for c in range(1, _ws2.max_column + 1)
          if "事业部" in str(_ws2.cell(row=1, column=c).value or "")
          or "事业部" in str(_ws2.cell(row=2, column=c).value or "")) == 1)
_wb2.close()
os.remove(TMP4)

# ---- 4.12 匹配结果：别名指向已有列时，预览直接填进那一列（不单列新列） ----
_alias_m = pd.DataFrame({"事业部": ["", ""], "项目": ["朗晴居", "景安花园"]})
_alias_k = pd.DataFrame({"项目名称": ["朗晴居", "景安花园"], "事业部": ["华二第三事业部", "华二第一事业部"]})
_alias_res = run_match(_alias_m, _alias_k, ["项目"], ["项目名称"], [], [],
                       ["事业部"], mode="fuzzy",
                       fetch_alias={"事业部": "事业部"})
_acols = list(_alias_res["result"].columns)
check("别名结果：不出现「事业部(钥匙表)」新列", "事业部(钥匙表)" not in _acols)
check("别名结果：列数=匹配表原列+3注释列+备注",
      _acols == ["事业部", "项目", "匹配轮次", "匹配率 (%)", "未匹配原因", "备注"])
check("别名结果：事业部列被就地填入补缺值",
      list(_alias_res["result"]["事业部"]) == ["华二第三事业部", "华二第一事业部"])

# ---- 4.13 经验库（兜底语义）：精确优先 / 模糊优先 / 无候选兜底 / 消歧 / 覆盖 / 忽略 ----
from core.experience import ExperienceStore
_exp_path = os.path.join(S, "_tmp_experience.json")
if os.path.exists(_exp_path):
    os.remove(_exp_path)
_exp = ExperienceStore(path=_exp_path)

# (1) 模糊命中：库为空也走模糊；记库后模糊仍优先（库只在配不上时兜底）
_em = pd.DataFrame({"项目": ["朗晴居一"]})
_ek = pd.DataFrame({"项目名称": ["朗晴居一期"], "事业部": ["华二第三事业部"]})
_r0 = run_match(_em, _ek, ["项目"], ["项目名称"], [], [], ["事业部"],
                mode="fuzzy", threshold=80, experience=_exp)
check("经验库兜底：库为空时走模糊", _r0["row_status"][0]["status"] == "钥匙模糊")
_exp.record("朗晴居一", "朗晴居一期", src="manual")
_r1 = run_match(_em, _ek, ["项目"], ["项目名称"], [], [], ["事业部"],
                mode="fuzzy", threshold=80, experience=_exp)
check("经验库兜底：模糊能配上时库不抢答（仍走模糊）", _r1["row_status"][0]["status"] == "钥匙模糊")

# (2) 精确优先于库（新之助→小爱 记录，但当前精确匹配到 野原）
_exp.record("新之助", "小爱", src="manual")
_km = pd.DataFrame({"项目": ["新之助"]})
_kk = pd.DataFrame({"项目名称": ["新之助", "小爱"], "事业部": ["野原", "小爱"]})
_r2 = run_match(_km, _kk, ["项目"], ["项目名称"], [], [], ["事业部"],
                mode="fuzzy", experience=_exp)
check("经验库兜底：精确优先于库（填野原，不填库里的 小爱）",
      _r2["row_status"][0]["status"] == "钥匙精确"
      and _r2["result"]["事业部"].iloc[0] == "野原")

# (3) 无候选时的兜底：库里没有能配上的规则 → 库给答案
_km3 = pd.DataFrame({"项目": ["新之助"]})
_kk3 = pd.DataFrame({"项目名称": ["小爱"], "事业部": ["小爱事业部"]})
_r3 = run_match(_km3, _kk3, ["项目"], ["项目名称"], [], [], ["事业部"],
                mode="fuzzy", threshold=80, experience=_exp)
check("经验库兜底：规则都配不上时库兜底命中",
      _r3["row_status"][0]["status"] == "经验库"
      and _r3["result"]["事业部"].iloc[0] == "小爱事业部")

# (4) 重复预警消歧：候选是"相似但不同"的键（新之助-野原 / 新之助-小爱），库里指定 小爱
_kk4 = pd.DataFrame({"项目名称": ["新之助-野原", "新之助-小爱"],
                     "事业部": ["野原事业部", "小爱事业部"]})
_exp.record("新之助", "新之助-小爱", src="manual")
_r4 = run_match(_km3, _kk4, ["项目"], ["项目名称"], [], [], ["事业部"],
                mode="fuzzy", threshold=50, experience=_exp)
check("经验库兜底：对上多条时按库消歧（选 新之助-小爱）",
      _r4["row_status"][0]["status"] == "经验库"
      and _r4["result"]["事业部"].iloc[0] == "小爱事业部")

# (5) 覆盖生效：同键再记别的 → 立即覆盖
_exp.record("新之助", "新之助-野原", src="manual")
check("经验库：同键再确认=覆盖（后确认者赢）", _exp.build_lookup().get("新之助") == "新之助-野原")
_exp.record("新之助", "新之助-小爱", src="manual")
check("经验库：覆盖可反复（再改回小爱）", _exp.build_lookup().get("新之助") == "新之助-小爱")

# (6) 忽略：规则配不上 → 直接跳过
_exp.record("不存在项目", ignore=True, src="manual")
_r6 = run_match(pd.DataFrame({"项目": ["不存在项目"]}), _ek, ["项目"], ["项目名称"],
                [], [], ["事业部"], mode="fuzzy", experience=_exp)
check("经验库：人工忽略生效", _r6["row_status"][0]["status"] == "人工忽略")

# (7) 备注列存在且标注来源
check("结果：含「备注」列且经验库兜底有说明",
      "备注" in _r3["result"].columns and "经验库兜底" in str(_r3["result"]["备注"].iloc[0]))
os.remove(_exp_path)


# ---- 4.14 报价单解析（M2a） ----
from core.parser_rule import parse_quote_file, supplier_from_filename
_qA = parse_quote_file(path=os.path.join(S, "报价单A_标准格式.xlsx"))[0]
check("解析：标准件表头行=1", _qA["header_row"] == 1)
check("解析：标准件字段映射含 品名/含税单价/税率",
      all(c in _qA["mapping"] for c in ("品名", "含税单价", "税率")))
check("解析：标准件行数=6", len(_qA["df"]) == 6)
check("解析：文件名提取供应商", supplier_from_filename("物资采购清单（利源通）(1).xls") == "利源通")
_qB = parse_quote_file(path=os.path.join(S, "报价单B_表头第2行_仅不含税.xlsx"))[0]
check("解析：变体件表头行=2 且识别不含税单价", _qB["header_row"] == 2 and "不含税单价" in _qB["mapping"])

# ---- 4.15 跨供应商对齐（M2b） ----
from core.aligner import align_quotes
_qs = [parse_quote_file(path=os.path.join(S, "报价单A_标准格式.xlsx"))[0],
       parse_quote_file(path=os.path.join(S, "报价单D_同品异名_缺行多行.xlsx"))[0]]
_ar = align_quotes(_qs, price_field="不含税单价")
check("对齐：A/D 两家 → 行数=7（6 品 + 劳保）", len(_ar["matrix"]) == 7)
check("对齐：含最低价与最低价供应商列",
      "最低价" in _ar["matrix"].columns and "最低价供应商" in _ar["matrix"].columns)
check("对齐：D 同品异名未拆行（螺钉 1 行）",
      sum(1 for n in _ar["matrix"]["品名"] if "螺钉" in str(n)) == 1)

# ---- 4.16 标红配色 + 规则建议（M2c） ----
from core.highlighter import PRESETS
from core.advisor_rule import advise
check("配色：浅色彩虹 ≥8 档且含浅黄", len(PRESETS) >= 8 and any(p[0] == "浅黄" for p in PRESETS))
_adm = pd.DataFrame({"品名": ["x"], "规格": [""], "甲": [100.0], "乙": [110.0], "丙": [5.0]})
_adv = advise(_adm, ["甲", "乙", "丙"], low_ratio=0.5)
check("建议：异常低价预警生效", _adv["rows"].iloc[0]["预警"] == "异常低价")
check("建议：汇总结构完整", {"品类数", "异常低价", "有缺报价的品类", "供应商数"} <= set(_adv["summary"]))

# ---- 4.17 LLM 语义兜底（离线：假客户端；不发网络）----
class _FakeLLM:
    enabled = True

    def __init__(self, payload):
        self.payload = payload
        self.calls = 0

    def chat(self, messages, **kw):
        self.calls += 1
        return {"text": self.payload, "usage": {}, "cost_usd": 0.0,
                "cached": False, "error": None, "ms": 1}


_mm = pd.DataFrame({"项目": ["紧固件"]})
_kkx = pd.DataFrame({"项目名称": ["螺钉M4×10"], "事业部": ["五金事业部"]})
_fl = _FakeLLM('{"matches":[{"row":0,"key":0,"confidence":0.9,"reason":"语义等价"}]}')
_rl = run_match(_mm, _kkx, ["项目"], ["项目名称"], [], [], ["事业部"],
                mode="fuzzy", threshold=80, llm=_fl)
check("LLM兜底：规则配不上时命中并标来源", _rl["row_status"][0]["status"] == "LLM语义")
check("LLM兜底：补缺值正确带入", _rl["result"]["事业部"].iloc[0] == "五金事业部")
check("LLM兜底：备注含置信度说明", "LLM兜底" in str(_rl["result"]["备注"].iloc[0]))
check("LLM兜底：统计含 LLM 计数", _rl["stats"].get("LLM语义") == 1)
_fl2 = _FakeLLM('{"matches":[{"row":0,"key":0,"confidence":0.9,"reason":"x"}]}')
_fl2.enabled = False
run_match(_mm, _kkx, ["项目"], ["项目名称"], [], [], ["事业部"], mode="fuzzy", llm=_fl2)
check("LLM兜底：未启用则不调用", _fl2.calls == 0)
_fl3 = _FakeLLM('{"matches":[{"row":0,"key":0,"confidence":0.3,"reason":"低置信"}]}')
_rl3 = run_match(_mm, _kkx, ["项目"], ["项目名称"], [], [], ["事业部"],
                 mode="fuzzy", threshold=80, llm=_fl3)
check("LLM兜底：低于阈值视为未匹配", _rl3["row_status"][0]["status"] == "未匹配")

# ---- 4.18 非标解析（M3c）：规则优先 / 非标走 LLM（离线假客户端） ----
from core.parser_llm import parse_quote_auto as _pqa
TMP5 = os.path.join(S, "_tmp_weird.xlsx")
_wb5 = _WB(); _ws5 = _wb5.active; _ws5.title = "报价"
_ws5.append(["料号", "描述", "计量单位", "需求数量", "采购价", "增值税"])
_ws5.append(["A001", "不锈钢螺丝M4", "个", 100, 1.13, "13%"])
_wb5.save(TMP5)


class _FLLM:
    enabled = True

    def __init__(self, p):
        self.p = p
        self.calls = 0

    def chat(self, messages, **kw):
        self.calls += 1
        return {"text": self.p, "usage": {}, "cost_usd": 0.0, "cached": False, "error": None, "ms": 1}


_fw = _FLLM('{"header_row":1,"mapping":{"品名":2,"单位":3,"数量":4,'
            '"含税单价":5,"税率":6},"supplier":"S"}')
_qw, _uw = _pqa(path=TMP5, llm=_fw)
check("非标解析：触发 LLM 且映射生效",
      _uw is True and _fw.calls == 1 and _qw[0]["df"]["品名"].iloc[0] == "不锈钢螺丝M4")
_fr = _FLLM('{}')
_qr, _ur = _pqa(path=os.path.join(S, "报价单A_标准格式.xlsx"), llm=_fr)
check("非标解析：标准件规则优先不调用 LLM", _ur is False and _fr.calls == 0)
os.remove(TMP5)

# ---- 5. 与 ground truth 核对（A 基准价=浮动 0） ----
import json
with open(os.path.join(S, "ground_truth.json"), encoding="utf-8") as f:
    truth = json.load(f)
for item in truth["items"]:
    expect = item["prices_tax_in"]["A"]
    actual = to_number(a_price[item["canonical"]])
    check(f"ground truth：{item['canonical']} A价={expect}", abs(actual - expect) < 0.01)

print(f"\n===== 冒烟测试全部通过：{ok} 项断言 =====")
