# -*- coding: utf-8 -*-
"""表格美化测试：合成表（格式断言）+ 真实文件「供应商信息反馈.xlsx」（可选高亮）。

运行：py tests/experiment_beautify.py
"""
import datetime as dt
import io
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

from openpyxl import Workbook, load_workbook

from core.theme import COLORS, beautify_bytes, estimate_width, guess_number_format

REAL = os.path.join(_ROOT, "data", "ground_truth", "供应商信息反馈.xlsx")
ok = 0


def check(name, cond):
    global ok
    assert cond, f"[FAIL] {name}"
    ok += 1
    print(f"[PASS] {name}")


def _sample_wb():
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["供应商名称很长很长用于测列宽", "工资（元）", "税率", "日期", "备注"])
    ws.append(["甲", 1234.5, 0.13, dt.date(2026, 9, 14), "ok"])
    ws.append(["乙", 300, 0.01, dt.date(2026, 1, 1), "x"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


# ---- 工具函数 ----
check("数字格式：金额→两位小数", guess_number_format("工资（元）", [1.5, 2]) == "#,##0.00")
check("数字格式：整数→千分位无小数", guess_number_format("数量", [1, 2]) == "#,##0")
check("数字格式：比例→百分比", guess_number_format("税率", [0.13, 0.01]) == "0.00%")
check("数字格式：日期列→yyyy-mm-dd", guess_number_format("日期", [dt.date(2026, 1, 1)]) == "yyyy-mm-dd")
check("列宽：中文按 2 字符估算且受上下限约束",
      8 <= estimate_width(["短"], "供应商名称很长很长用于测列宽", 8, 40) <= 40)

# ---- 美化（新增 sheet，原表不动） ----
data = _sample_wb()
out, info = beautify_bytes(data, "sample.xlsx", "明细", 1, {"banded": True})
wb = load_workbook(io.BytesIO(out))
check("美化：新增 Sheet「美化-明细」且原 sheet 保留", len(wb.sheetnames) == 2
      and wb.sheetnames[1].startswith("美化-") and "明细" in wb.sheetnames)
orig, new = wb["明细"], wb.worksheets[1]
check("美化：原 Sheet 未被改动（无冻结/无表头色）",
      orig.freeze_panes is None and orig.cell(row=1, column=1).fill.patternType is None)
check("美化：表头深底白字+加粗",
      new.cell(row=1, column=1).fill.fgColor.rgb == "FF44546A"
      and new.cell(row=1, column=1).font.color.rgb == "FFFFFFFF"
      and new.cell(row=1, column=1).font.bold is True)
check("美化：冻结到数据首行 + 自动筛选覆盖数据区",
      new.freeze_panes == "A2" and str(new.auto_filter.ref) == "A1:E3")
check("美化：数字格式已套用（工资两位小数、税率百分比、日期）",
      new["B2"].number_format == "#,##0.00" and new["C2"].number_format == "0.00%"
      and new["D2"].number_format == "yyyy-mm-dd")
check("美化：列宽自适应（不再固定 18）",
      abs(new.column_dimensions["A"].width - 18) > 0.5)
check("美化：表下写了页脚（生成时间）",
      any("生成时间" in str(new.cell(row=r, column=1).value or "")
          for r in range(new.max_row - 4, new.max_row + 1)))
wb.close()

# ---- 最低值高亮（可选）+ 真实文件 ----
if os.path.exists(REAL):
    with open(REAL, "rb") as f:
        real = f.read()
    out2, info2 = beautify_bytes(real, os.path.basename(REAL), "Sheet1", 2, {"highlight_min": True})
    wb2 = load_workbook(io.BytesIO(out2))
    ws2 = wb2.worksheets[1]
    b4 = ws2.cell(row=4, column=2).fill.fgColor.rgb.upper()
    c4 = ws2.cell(row=4, column=4).fill.fgColor.rgb.upper()
    check("真实文件：表头行(第2行)已美化", ws2.cell(row=2, column=1).fill.fgColor.rgb == "FF44546A")
    check("真实文件：冻结在数据首行(A3)", ws2.freeze_panes == "A3")
    check("真实文件：工资列最低值高亮=结论绿",
          b4.endswith(COLORS["conclusion"][0].lstrip("#")))
    check("真实文件：不可比列（发薪时间/月）未被高亮", not c4.endswith(COLORS["conclusion"][0].lstrip("#")))
    check("真实文件：原 Sheet 数据未被改动",
          wb2["Sheet1"].cell(row=4, column=2).value == ws2.cell(row=4, column=2).value)
    wb2.close()
else:
    check("真实文件：data/ground_truth/供应商信息反馈.xlsx 存在", False)

# ---- 就地美化（in_place） ----
out3, info3 = beautify_bytes(data, "sample.xlsx", "明细", 1, {"in_place": True})
wb3 = load_workbook(io.BytesIO(out3))
check("就地美化：Sheet 名不变且已套表头样式",
      len(wb3.sheetnames) == 1 and wb3["明细"].cell(row=1, column=1).fill.fgColor.rgb == "FF44546A")
check("就地美化：数据默认居中", wb3["明细"].cell(row=2, column=1).alignment.horizontal == "center")
wb3.close()

# ---- 就地美化 + 备份（回归：备份不能被旧内存字节覆盖） ----
import shutil
from core.theme import beautify_file_in_place
TMPF = os.path.join(os.environ["TEMP"], "opencode", "beautify_inplace.xlsx")
os.makedirs(os.path.dirname(TMPF), exist_ok=True)
shutil.copyfile(REAL, TMPF)
bk1, _i1 = beautify_file_in_place(TMPF, "Sheet1", 2, {})
wb4 = load_workbook(TMPF)
check("就地美化：备份 sheet 与美化后 sheet 同时存在",
      bk1.startswith("ai副本") and bk1 in wb4.sheetnames and "Sheet1" in wb4.sheetnames)
check("就地美化：Sheet1 已美化且数据未变",
      wb4["Sheet1"].cell(row=2, column=1).fill.fgColor.rgb == "FF44546A"
      and wb4["Sheet1"].cell(row=4, column=2).value == 300)
check("就地美化：备份内容=美化前（表头底色≠美化后的深蓝）",
      str(wb4[bk1].cell(row=2, column=1).fill.fgColor.rgb) != "FF44546A")
wb4.close()
bk2, _i2 = beautify_file_in_place(TMPF, "Sheet1", 2, {})
wb5 = load_workbook(TMPF)
check("就地美化：二次执行备份递增且旧备份仍在",
      bk2 != bk1 and bk1 in wb5.sheetnames and bk2 in wb5.sheetnames)
wb5.close()
os.remove(TMPF)

print(f"\n===== 表格美化测试通过：{ok} 项断言 =====")
