# -*- coding: utf-8 -*-
"""真实文件实验：两表匹配 · 合并单元格链路（临时副本上跑，不碰原文件）。

匹配表 = 001.xlsx · 合同进度表（两行合并表头：事业部/项目/类别竖向合并 +
「合同信息」横向合并跨 3 个月份列 + 数据区类别竖向合并）。
钥匙表 = 采购项目汇总-2026年12月.xlsx · Sheet1（表头竖向合并 A1:A2..F1:F2）。
验证链路：顶行自动降级 → 拆分命名（前缀回填）→ 数据区合并填充 → run_match → 写回临时副本。
运行：py tests/experiment_pipei.py
"""
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8")
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

from core.excel_io import (backup_sheet_numbered, detect_header_block_bottom,
                           read_table, split_two_row_header, unmerge_fill_data,
                           write_result_in_sheet)
from core.matcher import run_match

# 项目内基准文件（不再依赖桌面路径）
M_REAL = os.path.join(_ROOT, "data", "ground_truth", "匹配表_001_合同进度表.xlsx")
K_REAL = os.path.join(_ROOT, "data", "ground_truth", "钥匙表_采购项目汇总.xlsx")
M_SHEET = "ai副本合同进度表"   # 写回前的原始快照（不受用户验收写回影响）
K_SHEET = "Sheet1"
TMP_M = os.path.join(os.environ["TEMP"], "001_实验副本.xlsx")
TMP_K = os.path.join(os.environ["TEMP"], "采购项目汇总_实验副本.xlsx")
shutil.copyfile(M_REAL, TMP_M)
shutil.copyfile(K_REAL, TMP_K)

ok = 0


def check(name, cond):
    global ok
    assert cond, f"[FAIL] {name}"
    ok += 1
    print(f"[PASS] {name}")


def load_flow(path, sheet, hdr_pick):
    """模拟 app.load_table 的读取流程（自动降级 + 拆分 + 填充）。"""
    with open(path, "rb") as f:
        fb = f.read()
    bottom = detect_header_block_bottom(fb, path, sheet, hdr_pick)
    hdr = bottom or hdr_pick
    names, is_two = split_two_row_header(fb, path, sheet, hdr)
    df = read_table(path, sheet, header_row=hdr)
    if names and is_two:
        df = df.rename(columns={df.columns[j]: names[j]
                                for j in range(min(len(names), len(df.columns)))})
    df, n_fill = unmerge_fill_data(fb, path, sheet, hdr, df)
    return df, hdr, bottom, names, is_two, n_fill


# ---- 1. 匹配表：顶行降级 + 拆分 + 填充 ----
m_df, m_hdr, m_bot, m_names, m_is_two, m_fill = load_flow(TMP_M, M_SHEET, 1)
check("匹配表：第1行被识别为两行表头顶部 → 降级到第2行", m_bot == 2 and m_hdr == 2)
check("匹配表：横向合并前缀回填（合同信息 2026年11月）",
      "合同信息 2026年11月" in list(m_df.columns)
      and "合同信息 2026年12月" in list(m_df.columns))
check("匹配表：竖向合并列保持原名（事业部/项目/类别）",
      all(c in list(m_df.columns) for c in ("事业部", "项目", "类别", "责任人", "合同进度")))
check("匹配表：数据区合并填充 >0 格", m_fill > 0)
check("匹配表：类别列每行都有值（消防维保等不再断档）",
      m_df["类别"].astype(str).str.strip().ne("").all())
check("匹配表：行数不变", len(m_df) == 37)

# ---- 2. 钥匙表：竖向合并表头降级 + 命名 ----
k_df, k_hdr, k_bot, k_names, k_is_two, k_fill = load_flow(TMP_K, K_SHEET, 1)
check("钥匙表：第1行被识别为表头顶部 → 降级到第2行", k_bot == 2 and k_hdr == 2)
check("钥匙表：列名完整（项目名称…合同编号，无 Unnamed）",
      list(k_df.columns)[:6] == ["项目名称", "采购三级分类", "起始日期", "截止日期",
                                 "合同到期月份", "合同编号"])
check("钥匙表：数据首行是真实数据（不是表头副本）",
      str(k_df["项目名称"].iloc[0]).strip() != "项目名称")

# ---- 3. run_match：项目 ↔ 项目名称 ----
logs = []
res = run_match(m_df, k_df, ["项目"], ["项目名称"], [], [],
                ["合同编号", "合同到期月份"], mode="fuzzy", threshold=80.0,
                log=logs.append)
stats = res["stats"]
check("匹配：朗晴居二期/一期/怡安花园/景安花园 精确命中（≥4）",
      stats["钥匙精确"] >= 4)
check("匹配：结果行数 ≥ 匹配表行数（重复钥匙按组合展开）",
      len(res["result"]) >= len(m_df))
row_lq = res["result"][res["result"]["项目"].astype(str) == "朗晴居二期"]
check("匹配：朗晴居二期 补到了合同编号",
      len(row_lq) >= 1
      and any(str(v).startswith("YC-") for v in row_lq["合同编号"]))

# ---- 4. 写回临时副本（备份 + 就地写结果） ----
bk = backup_sheet_numbered(TMP_M, M_SHEET)
check("备份：生成 ai副本 Sheet", bk and bk.startswith("ai副本"))
row_map = list(range(len(m_df)))
written, overwritten = write_result_in_sheet(
    TMP_M, M_SHEET, m_hdr, row_map, res["row_status"], res["fetch_records"],
    ["合同编号", "合同到期月份"])
check("写回：返回写入列", "合同编号" in written and "合同到期月份" in written)

from openpyxl import load_workbook
wb = load_workbook(TMP_M)
ws = wb[M_SHEET]
headers = [str(ws.cell(row=m_hdr, column=c).value or "") for c in range(1, ws.max_column + 1)]
i_no = headers.index("合同编号") + 1
i_month = headers.index("合同到期月份") + 1
check("写回：朗晴居二期行补上合同编号",
      any(str(ws.cell(row=m_hdr + 1 + i, column=i_no).value or "").startswith("YC-")
          for i in m_df.index[m_df["项目"].astype(str) == "朗晴居二期"]))
check("写回：注释列（匹配轮次/匹配率/未匹配原因）已追加",
      all(c in headers for c in ("匹配轮次", "未匹配原因")))
wb.close()

os.remove(TMP_M)
os.remove(TMP_K)
print(f"\n===== 两表匹配真实文件实验（001.xlsx ↔ 采购项目汇总·临时副本）全部通过：{ok} 项断言 =====")
print("原文件未做任何改动；你可在网页里对这两个文件实际操作验收。")
